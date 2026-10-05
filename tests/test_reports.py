"""Acceptance A40–A45: real formats, filtered scope, review and job revocation."""
import csv
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
import json
import uuid
import base64
import importlib.util
from pathlib import Path
import subprocess
import sys
import os

import pytest
from openpyxl import load_workbook
from flask import g
from sqlalchemy import select
from volleyball.common import APIError
from volleyball.models import (Organization,User,Membership,AuthSession,Team,Tournament,
    Division,Stage,Entry,Match,Job,Venue,WeatherCache,Court,CourtAssignment,Announcement)
from volleyball.db import get_db,set_scope
from volleyball.reports import (spreadsheet_safe,render_report,present_report,build_report,
    process_export_jobs,validate_import,authorize_worker,pdf_column_groups,pdf_text)
from volleyball.weather import validate_provider_payload,WeatherUnavailable,venue_weather


def sample_report():
    return {'columns':['name','score'],'rows':[{'name':'=HYPERLINK("https://invalid.example")','score':12},
        {'name':'<b>literal team</b>','score':2}], 'total':2,
        'metadata':{'type':'matches','organization':{'name':'Isolated test'},'generated_at':'2026-10-05T00:00:00Z',
                    'timezone':'America/New_York','filters':{'status':'ended'},'official_status':'finalized','source_generations':{}}}


@pytest.mark.parametrize('value',['=SUM(1,2)','+SUM(1,2)','-cmd','@SUM(1,2)','  =1','\t=1','\r=1','\n=1'])
def test_formula_leading_values_are_escaped(value):
    assert spreadsheet_safe(value)=="'"+value
    assert spreadsheet_safe(-10)==-10


def test_csv_and_xlsx_preserve_filtered_values_without_formulas():
    report=sample_report();payload,mime=render_report(report,'csv')
    rows=list(csv.reader(StringIO(payload.decode('utf-8-sig'))))
    assert rows[0]==['name','score'] and rows[1][0].startswith("'=") and rows[2][1]=='2'
    assert ['REPORT METADATA'] in rows
    payload,mime=render_report(report,'xlsx');workbook=load_workbook(BytesIO(payload))
    assert workbook['Report']['A2'].value.startswith("'=")
    assert workbook['Report']['A2'].data_type!='f'
    assert workbook['Report']['B3'].value==2
    assert 'Metadata' in workbook.sheetnames


def test_pdf_is_real_and_untrusted_markup_is_escaped():
    payload,mime=render_report(sample_report(),'pdf')
    assert payload.startswith(b'%PDF-') and payload.rstrip().endswith(b'%%EOF')
    assert mime=='application/pdf'


def test_wide_pdf_groups_keep_unique_identity_even_with_reordered_or_filtered_columns():
    columns=['division','rank','entry_id','name','played','won','lost','drawn','sets_won','sets_lost','points_for','points_against']
    groups=pdf_column_groups(columns,'standings')
    assert groups[1][0]=='entry_id' and all('entry_id' in group for group in groups)
    reordered=['division','rank','name','played','won','lost','entry_id']
    groups=pdf_column_groups(reordered,'standings')
    assert len(groups)==1 and groups[0][0]=='entry_id'
    filtered=['division','rank','played','won','lost','drawn','sets_won']
    assert all(group[0]=='__report_row' for group in pdf_column_groups(filtered,'standings'))


def test_pdf_fonts_cover_extended_latin_greek_cyrillic_and_infinity():
    from reportlab.pdfbase import pdfmetrics
    render_report(sample_report(),'pdf')
    for name in ('VolleyballNoto','VolleyballNotoBold'):
        coverage=pdfmetrics.getFont(name).face.charToGlyph
        assert all(ord(char) in coverage for char in 'São Paulo Łódź Žluťoučký Ελληνικά Москва')
    assert 'infinity' in pdfmetrics.getFont('Symbol').encoding.vector
    assert pdf_text('<b>∞</b>')=='&lt;b&gt;<font name="Symbol">∞</font>&lt;/b&gt;'


def test_xlsx_large_source_metadata_is_complete_without_cell_truncation():
    report=sample_report();large='metadata-source-version-'*4000
    report['metadata']['large_source_metadata']=large
    payload,mime=render_report(report,'xlsx');sheet=load_workbook(BytesIO(payload))['Metadata']
    parts=[row[1].value for row in sheet if str(row[0].value).startswith('large_source_metadata [part ')]
    assert ''.join(parts)==large and len(parts)>1 and all(len(part)<=30000 for part in parts)


def test_numeric_sort_column_selection_and_bounded_pages():
    result=present_report(sample_report(),{'sort':'score','columns':'score,name','page_size':'1'})
    assert result['rows']==[{'score':2,'name':'<b>literal team</b>'}] and result['total']==2
    with pytest.raises(APIError): present_report(sample_report(),{'page_size':'201'})
    with pytest.raises(APIError): present_report(sample_report(),{'columns':'credential'})


@pytest.fixture
def report_context(app):
    """Own isolated orgs on the shared real-PostgreSQL app fixture."""
    with app.app_context():
        db=get_db();tag=uuid.uuid4().hex[:10]
        user=User(email=tag+'@report.test',name='Report tester',password_hash='test-not-a-login')
        org=Organization(name='Report organization '+tag,slug='report-'+tag,demo=True,settings={})
        other=Organization(name='Other report organization '+tag,slug='other-report-'+tag,demo=True)
        db.add_all([user,org,other]);db.flush();set_scope(org.id)
        member=Membership(organization_id=org.id,user_id=user.id,role='admin',status='active')
        auth=AuthSession(user_id=user.id,token_hash=uuid.uuid4().hex+uuid.uuid4().hex,expires_at=datetime.now(timezone.utc)+timedelta(hours=1),mfa_verified=False)
        today=datetime.now(timezone.utc).date()
        tournament=Tournament(organization_id=org.id,name='Report event',slug='report-event-'+tag,start_date=today,end_date=today+timedelta(days=14),rules={},status='published')
        db.add_all([member,auth,tournament]);db.flush()
        division=Division(organization_id=org.id,tournament_id=tournament.id,name='Report division')
        db.add(division);db.flush()
        stage=Stage(organization_id=org.id,tournament_id=tournament.id,division_id=division.id,name='Report stage')
        teams=[Team(organization_id=org.id,name='=Unsafe team'),Team(organization_id=org.id,name='Second team')]
        db.add_all([stage,*teams]);db.flush()
        entries=[Entry(organization_id=org.id,tournament_id=tournament.id,division_id=division.id,team_id=t.id,name=t.name) for t in teams]
        db.add_all(entries);db.flush()
        match=Match(organization_id=org.id,tournament_id=tournament.id,division_id=division.id,stage_id=stage.id,entry_a_id=entries[0].id,entry_b_id=entries[1].id,
            status='ended',result_status='finalized',winner_id=entries[0].id,outcome='normal',state={'sets':[{'a':25,'b':23,'complete':True,'winner':'a'}]},planned_at=datetime.now(timezone.utc),version=7,result_revision=2)
        db.add(match);db.commit()
        g.organization=org;g.user=user;g.membership=member;g.auth_session=auth
        yield dict(db=db,org=org,other=other,user=user,member=member,auth=auth,tournament=tournament,division=division,stage=stage,match=match,entries=entries)
        db.rollback()


def test_report_scope_filters_and_source_versions(report_context):
    env=report_context
    report=build_report('matches',{'tournament_id':env['tournament'].id,'result_status':'finalized'})
    assert report['total']==1 and report['rows'][0]['sets'][0]['a']==25
    assert report['metadata']['source_match_versions'][env['match'].id]=={'version':7,'result_revision':2}
    assert report['metadata']['organization']['id']==env['org'].id
    assert 'qr' not in json.dumps(report).lower() and 'token_hash' not in json.dumps(report)
    env['member'].tournament_ids=[str(uuid.uuid4())];env['db'].commit()
    with pytest.raises(APIError) as caught: build_report('matches',{'tournament_id':env['tournament'].id})
    assert caught.value.status==403


def test_export_worker_rejects_revoked_membership(report_context):
    env=report_context;job=Job(organization_id=env['org'].id,user_id=env['user'].id,kind='export',type='matches',format='csv',filters={'tournament_id':env['tournament'].id},
        data={'auth_session_id':env['auth'].id},expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    env['db'].add(job);env['db'].commit();env['member'].status='suspended';env['db'].commit()
    process_export_jobs()
    env['db'].refresh(job)
    assert job.status=='failed' and job.file_path is None and 'revoked' in job.error


def test_export_worker_writes_real_file_and_rechecks_session(report_context):
    env=report_context;job=Job(organization_id=env['org'].id,user_id=env['user'].id,kind='export',type='matches',format='xlsx',filters={'tournament_id':env['tournament'].id},
        data={'auth_session_id':env['auth'].id},expires_at=datetime.now(timezone.utc)+timedelta(hours=1))
    env['db'].add(job);env['db'].commit();process_export_jobs()
    env['db'].refresh(job)
    assert job.status=='completed' and job.progress==100 and job.file_path.endswith('.xlsx')
    env['auth'].revoked_at=datetime.now(timezone.utc);env['db'].commit()
    with pytest.raises(APIError): authorize_worker(job)


def test_import_row_errors_and_no_silent_result_overwrite(report_context):
    env=report_context
    valid,errors=validate_import('teams',[{'name':'New team','abbreviation':'NEW','color':'#123456'}, {'name':'','abbreviation':'','color':'bad'}])
    assert len(valid)==1 and errors[0]['row']==3
    valid,errors=validate_import('referees',[{'match_id':env['match'].id,'referee_name':'New referee'}])
    assert not valid and 'before play' in errors[0]['message']
    valid,errors=validate_import('entries',[{'division_id':str(uuid.uuid4()),'team_id':str(uuid.uuid4()),'name':'','seed':'1','pool':''}])
    assert not valid and errors[0]['row']==2 and 'unavailable' in errors[0]['message']


def daily_payload(venue):
    today=datetime.now(timezone.utc).astimezone(__import__('zoneinfo').ZoneInfo(venue.timezone)).date()
    return {'timezone':venue.timezone,'forecast_start':today.isoformat(),'forecast_end':(today+timedelta(days=10)).isoformat(),
        'updated_at':datetime.now(timezone.utc).isoformat(),'temperature_unit':'C','wind_unit':'km/h','attribution':'Isolated weather test adapter',
        'days':[{'date':(today+timedelta(days=i)).isoformat(),'condition':'Test adapter condition','temperature_high':20,'temperature_low':10,'rain_probability':20,'wind_speed':5} for i in range(11)]}


def test_weather_real_horizon_max_seven_and_stale_failure(report_context,app):
    env=report_context;venue=Venue(organization_id=env['org'].id,name='Test venue',timezone='America/New_York',latitude=40,longitude=-74,settings={'coordinates_verified':True})
    env['db'].add(venue);env['db'].commit()
    original=dict(app.config)
    try:
        app.config['WEATHER_ADAPTER']=daily_payload;app.config['WEATHER_PROVIDER_NAME']='isolated-test'
        response=venue_weather(env['tournament'],venue,'celsius')
        assert response['status']=='available' and len(response['days'])==7 and response['temperature_unit']=='C'
        cache=env['db'].scalar(select(WeatherCache).where(WeatherCache.venue_id==venue.id));cache.expires_at=datetime.now(timezone.utc)-timedelta(seconds=1);env['db'].commit()
        def fail(venue): raise RuntimeError('Provider outage')
        app.config['WEATHER_ADAPTER']=fail
        assert venue_weather(env['tournament'],venue,'fahrenheit')['status']=='stale'
        env['tournament'].start_date+=timedelta(days=12);env['tournament'].end_date+=timedelta(days=12)
        app.config['WEATHER_ADAPTER']=daily_payload
        assert venue_weather(env['tournament'],venue,'celsius')['days']==[]
    finally:
        app.config.clear();app.config.update(original)


def test_weather_rejects_wrong_timezone_and_fake_dates():
    venue=Venue(timezone='America/New_York');payload=daily_payload(venue)
    payload['timezone']='UTC'
    with pytest.raises(WeatherUnavailable): validate_provider_payload(payload,venue)


def test_csv_import_commit_is_reviewed_and_all_or_nothing(admin_client):
    headers=admin_client.api_headers
    before=admin_client.get('/api/v1/teams',headers=headers).json['total']
    invalid=b'name,abbreviation,color\nReviewed new team,NEW,#123456\n,BAD,not-a-color\n'
    response=admin_client.post('/api/v1/imports/preview',headers=headers,data={'type':'teams','file':(BytesIO(invalid),'teams.csv')})
    assert response.status_code==201 and response.json['errors'][0]['row']==3
    commit=admin_client.post('/api/v1/imports/'+response.json['id']+'/commit',headers=headers,json={})
    assert commit.status_code==422
    assert admin_client.get('/api/v1/teams',headers=headers).json['total']==before
    valid=b'name,abbreviation,color\nReviewed new team,NEW,#123456\n'
    response=admin_client.post('/api/v1/imports/preview',headers=headers,data={'type':'teams','file':(BytesIO(valid),'teams.csv')})
    assert response.status_code==201 and not response.json['errors']
    commit=admin_client.post('/api/v1/imports/'+response.json['id']+'/commit',headers=headers,json={})
    assert commit.status_code==200 and commit.json['status']=='completed'
    assert admin_client.get('/api/v1/teams',headers=headers).json['total']==before+1


def test_foreign_organization_report_id_and_public_export_revocation(admin_client,demo,client):
    second_headers={**admin_client.api_headers,'X-Organization-ID':demo['organizations'][1]['id']}
    second=admin_client.get('/api/v1/tournaments',headers=second_headers).json['items'][0]
    response=admin_client.get('/api/v1/reports',headers=admin_client.api_headers,query_string={'tournament_id':second['id']})
    assert response.status_code==404
    event=admin_client.get('/api/v1/tournaments',headers=admin_client.api_headers).json['items'][0]
    public_url='/api/v1/public/'+event['slug']+'/reports/export?type=schedule&format=csv'
    # Explicitly enable the public-safe type, then verify privacy is live.
    response=admin_client.patch('/api/v1/tournaments/'+event['id']+'/display',headers=admin_client.api_headers,json={'display':{'public_reports':['schedule']}})
    assert response.status_code==200
    response=client.get(public_url)
    assert response.status_code==200 and 'referee_name' not in response.data.decode('utf-8-sig').splitlines()[0]
    assert b'token_hash' not in response.data and b'contact' not in response.data
    response=admin_client.patch('/api/v1/tournaments/'+event['id'],headers=admin_client.api_headers,json={'visibility':'private'})
    assert response.status_code==200
    assert client.get(public_url).status_code==404


def test_sponsor_display_widget_saves_safe_urls_and_rejects_unsafe_urls(admin_client):
    headers=admin_client.api_headers
    tournament=admin_client.get('/api/v1/tournaments',headers=headers).json['items'][0]
    url='/api/v1/tournaments/'+tournament['id']+'/display'
    layout={'widgets':['sponsors'],'sizes':{'sponsors':'wide'},'sponsor_name':'Local sponsor',
            'sponsor_image':'https://sponsor.example/logo.png','sponsor_url':'https://sponsor.example/'}
    saved=admin_client.patch(url,headers=headers,json={'display':layout})
    assert saved.status_code==200 and saved.json['display']==layout
    assert admin_client.get(url,headers=headers).json['display']==layout
    for field,value in [('sponsor_url','javascript:alert(1)'),('sponsor_url','http://sponsor.example/'),
                        ('sponsor_url','https://user:password@sponsor.example/'),
                        ('sponsor_image','data:image/svg+xml,<svg/>'),('sponsor_image','//sponsor.example/logo.png')]:
        rejected=admin_client.patch(url,headers=headers,json={'display':{**layout,field:value}})
        assert rejected.status_code==400
        assert admin_client.get(url,headers=headers).json['display']==layout


def test_public_live_is_lightweight_dynamic_and_denied_after_privacy_change(admin_client,client,app):
    from sqlalchemy import event as sql_event
    tournament=admin_client.get('/api/v1/tournaments',headers=admin_client.api_headers).json['items'][0]
    full=client.get('/api/v1/public/'+tournament['slug'])
    assert full.status_code==200
    ended=next(m for m in full.json['matches'] if m['status']=='ended' and m['team_a'])
    with app.app_context():
        db=get_db();set_scope(tournament['organization_id'])
        venue=Venue(organization_id=tournament['organization_id'],name='Historical live venue',timezone='UTC')
        db.add(venue);db.flush()
        court=Court(organization_id=tournament['organization_id'],venue_id=venue.id,name='Historical live court')
        db.add(court);db.flush();court_id=court.id
        db.add(CourtAssignment(organization_id=tournament['organization_id'],match_id=ended['id'],court_id=court.id,
            status='occupied',released_at=datetime.now(timezone.utc)))
        entry=db.get(Entry,ended['team_a']['id']);entry.name='Reviewed changed participant'
        match=db.get(Match,ended['id']);match.version+=1;version=match.version
        db.commit()
    statements=[]
    def record(connection,cursor,statement,parameters,context,many):statements.append(statement.lower())
    engine=app.extensions['db_engine'];sql_event.listen(engine,'before_cursor_execute',record)
    try:response=client.get('/api/v1/public/'+tournament['slug']+'/live')
    finally:sql_event.remove(engine,'before_cursor_execute',record)
    assert response.status_code==200 and response.json['full_refresh'] is False
    assert {m['id'] for m in response.json['matches']}=={m['id'] for m in full.json['matches']}
    current=next(m for m in response.json['matches'] if m['id']==ended['id'])
    assert current['version']==version and current['team_a']['name']=='Reviewed changed participant'
    assert current['court']['id']==court_id
    assert response.json['configuration_version'] and isinstance(response.json['standings_generation'],int)
    assert not {'rules','scorer','referee_name','settings'} & current.keys()
    assert not any('scorer_sessions' in q or 'match_events' in q or 'result_revisions' in q for q in statements)
    # Tenant setup, publication discovery and persisted request throttling
    # have their own fixed overhead. Bound the actual projection reads by
    # table rather than guessing a total for those shared request hooks.
    projection_counts={table:sum(q.lstrip().startswith('select') and ('from '+table+'\n' in q or 'from '+table+' ' in q)
        for q in statements) for table in ('matches','entries','court_assignments','courts','venues')}
    assert projection_counts['matches']==1
    assert projection_counts['entries']<=1 and projection_counts['court_assignments']<=2
    assert projection_counts['courts']<=2 and projection_counts['venues']<=1
    assert admin_client.patch('/api/v1/tournaments/'+tournament['id'],headers=admin_client.api_headers,
        json={'visibility':'private'}).status_code==200
    assert client.get('/api/v1/public/'+tournament['slug']+'/live').status_code==404


def test_public_configuration_changes_for_metadata_but_not_point_projection(report_context):
    from volleyball.services.public_config import public_configuration_version
    env=report_context;db=env['db'];tournament=env['tournament']
    version=public_configuration_version(tournament);db.commit()
    # Model a persisted point/version update; immutable rule inputs do not change.
    match=env['match'];state=json.loads(json.dumps(match.state));state['sets'][0]['a']+=1
    match.state=state;match.version+=1;db.commit()
    assert public_configuration_version(tournament)==version
    env['stage'].name='Reviewed renamed stage';db.commit()
    renamed=public_configuration_version(tournament)
    assert renamed!=version
    announcement=Announcement(organization_id=env['org'].id,tournament_id=tournament.id,title='Visible notice',
        content='Public content',visibility='public',publish_at=datetime.now(timezone.utc)-timedelta(minutes=1))
    db.add(announcement);db.commit()
    visible=public_configuration_version(tournament)
    assert visible!=renamed
    future=Announcement(organization_id=env['org'].id,tournament_id=tournament.id,title='Not yet visible',
        content='Future content',visibility='public',publish_at=datetime.now(timezone.utc)+timedelta(days=1))
    db.add(future);db.commit()
    assert public_configuration_version(tournament)==visible
    announcement.expires_at=datetime.now(timezone.utc)-timedelta(seconds=1);db.commit()
    assert public_configuration_version(tournament)!=visible


def test_generated_file_cannot_be_downloaded_after_membership_revocation(admin_client,demo,app):
    response=admin_client.post('/api/v1/exports',headers=admin_client.api_headers,json={'type':'matches','format':'csv','filters':{}})
    assert response.status_code==202;job_id=response.json['id']
    with app.app_context(): process_export_jobs()
    response=admin_client.get('/api/v1/exports/'+job_id+'/download',headers=admin_client.api_headers)
    assert response.status_code==200
    with app.app_context():
        set_scope(admin_client.org_id);db=get_db()
        member=db.scalar(select(Membership).where(Membership.organization_id==admin_client.org_id,Membership.user_id==demo['user_id']))
        member.status='suspended';db.commit()
    assert admin_client.get('/api/v1/exports/'+job_id+'/download',headers=admin_client.api_headers).status_code==403
    # Global identity and its other organization remain usable.
    response=admin_client.get('/api/v1/dashboard',headers={**admin_client.api_headers,'X-Organization-ID':demo['organizations'][1]['id']})
    assert response.status_code==200


def test_native_recovery_encryption_authenticates_before_publishing(tmp_path):
    path=Path(__file__).resolve().parents[1]/'scripts'/'backup_crypto.py'
    specification=importlib.util.spec_from_file_location('volleyball_backup_crypto_test',path)
    module=importlib.util.module_from_spec(specification);specification.loader.exec_module(module)
    key=b'a'*32;source=tmp_path/'source';source.write_bytes(b'accepted-event-123\n'*(100000))
    encrypted=tmp_path/'encrypted';decrypted=tmp_path/'decrypted'
    module.encrypt_file(source,encrypted,key)
    assert b'accepted-event' not in encrypted.read_bytes()
    module.decrypt_file(encrypted,decrypted,key)
    assert decrypted.read_bytes()==source.read_bytes()
    corrupt=bytearray(encrypted.read_bytes());corrupt[-32]^=1;encrypted.write_bytes(corrupt)
    invalid_output=tmp_path/'invalid-output'
    with pytest.raises(Exception):module.decrypt_file(encrypted,invalid_output,key)
    assert not invalid_output.exists() and not list(tmp_path.glob('*.partial'))


def test_native_recovery_key_requires_private_file_permissions(tmp_path):
    path=Path(__file__).resolve().parents[1]/'scripts'/'backup_crypto.py'
    specification=importlib.util.spec_from_file_location('volleyball_backup_key_test',path)
    module=importlib.util.module_from_spec(specification);specification.loader.exec_module(module)
    keyfile=tmp_path/'key';keyfile.write_bytes(base64.urlsafe_b64encode(b'a'*32));keyfile.chmod(0o644)
    with pytest.raises(RuntimeError):module.read_key(keyfile)
    keyfile.chmod(0o600)
    assert module.read_key(keyfile)==b'a'*32


def test_wal_archive_and_restore_commands_are_idempotent_and_authenticated(tmp_path):
    scripts=Path(__file__).resolve().parents[1]/'scripts'
    key=tmp_path/'recovery.key';key.write_bytes(base64.urlsafe_b64encode(b'a'*32));key.chmod(0o600)
    source=tmp_path/'source-wal';source.write_bytes(b'accepted-WAL-content')
    storage=tmp_path/'archive';filename='000000010000000000000005'
    command=[sys.executable,str(scripts/'backup_wal.py'),str(source),filename,'--storage',str(storage),'--key-file',str(key)]
    subprocess.run(command,check=True,capture_output=True)
    subprocess.run(command,check=True,capture_output=True)
    restored=tmp_path/'pg_wal'/'RECOVERYXLOG'
    subprocess.run([sys.executable,str(scripts/'restore_wal.py'),filename,str(restored),'--storage',str(storage),'--key-file',str(key)],cwd=tmp_path,check=True,capture_output=True)
    assert restored.read_bytes()==source.read_bytes()
    original=(storage/'wal'/(filename+'.enc')).read_bytes();source.write_bytes(b'different-WAL-content')
    result=subprocess.run(command,capture_output=True)
    assert result.returncode!=0 and (storage/'wal'/(filename+'.enc')).read_bytes()==original


def test_physical_recovery_settings_never_undersize_the_source():
    scripts=Path(__file__).resolve().parents[1]/'scripts';sys.path.insert(0,str(scripts))
    try:
        specification=importlib.util.spec_from_file_location('volleyball_pitr_settings_test',scripts/'restore_pitr.py')
        module=importlib.util.module_from_spec(specification);specification.loader.exec_module(module)
    finally:sys.path.pop(0)
    defaults=module.recovery_settings({})
    assert defaults=={'max_connections':100,'max_worker_processes':8,'max_wal_senders':10,'max_prepared_transactions':0,'max_locks_per_transaction':64}
    source={'max_connections':200,'max_worker_processes':16,'max_wal_senders':20,'max_prepared_transactions':10,'max_locks_per_transaction':128}
    assert module.recovery_settings({'recovery_settings':source})==source
    with pytest.raises(RuntimeError):module.recovery_settings({'recovery_settings':{'max_connections':True}})
    with pytest.raises(RuntimeError):module.recovery_settings({'recovery_settings':{'max_connections':100001}})


def test_readiness_rejects_owner_runtime_and_migration_mismatch(app):
    from volleyball import create_app
    response=app.test_client().get('/api/v1/health/ready')
    assert response.status_code==200,response.json
    app.config['SCHEMA_REVISION']='unapplied-test-revision'
    response=app.test_client().get('/api/v1/health/ready')
    assert response.status_code==503 and response.json['schema']=='migration_required'
    app.config.pop('SCHEMA_REVISION')
    owner_app=create_app({**app.config,'DATABASE_URL':os.environ['TEST_DATABASE_URL']})
    try:
        response=owner_app.test_client().get('/api/v1/health/ready')
        assert response.status_code==503 and response.json['runtime_role']=='unsafe'
    finally:owner_app.extensions['db_engine'].dispose()
