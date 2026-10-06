"""Scoped persisted reporting, reviewed CSV imports and authorized export jobs."""
import csv
from datetime import datetime, date, time, timedelta, timezone
from html import escape
from io import BytesIO, StringIO
import json
import math
import re
from pathlib import Path
import uuid
from zoneinfo import ZoneInfo

import click
from flask import Blueprint, current_app, g, request, send_file
from sqlalchemy import select, or_, func
from sqlalchemy.orm import Session
from .auth import require_org
from .common import APIError, audit, body, iso, now, serialize
from .db import get_db, set_scope
from .models import (Organization, User, AuthSession, Membership, Team, Venue, Court, Tournament,
    Division, Stage, Entry, Match, CourtAssignment, MatchEvent, ResultRevision,
    Audit, Preference, Announcement, Job, RuleVersion, TournamentVenue)

bp = Blueprint('reports', __name__, url_prefix='/api/v1')
TYPES = {'summary', 'matches', 'standings', 'team_performance', 'schedule', 'court_use',
         'referees', 'outstanding_results', 'audit', 'history'}
PUBLIC_TYPES = {'summary', 'matches', 'standings', 'team_performance', 'schedule'}
FORMATS = {'csv', 'xlsx', 'pdf'}
FILTERS = {'tournament_id','division_id','stage_id','team_id','venue_id','court_id',
           'status','result_status','outcome','from','to','search'}
TEMPLATES = {
    'teams': ['name','abbreviation','color'],
    'entries': ['division_id','team_id','name','seed','pool'],
    'fixtures': ['stage_id','entry_a_id','entry_b_id','planned_at','planned_duration','referee_name'],
    'referees': ['match_id','referee_name'],
}
MAX_REPORT_ROWS = 50000
MAX_IMPORT_ROWS = 2000


def without_secrets(value):
    if isinstance(value,dict):
        return {key:without_secrets(child) for key,child in value.items()
                if not any(term in str(key).casefold() for term in ('password','secret','token','credential','csrf','qr_url','grant_url','backup_key'))}
    if isinstance(value,list): return [without_secrets(child) for child in value]
    return value


def ident(value, label='identifier'):
    try: return str(uuid.UUID(str(value)))
    except (ValueError, TypeError, AttributeError): raise APIError(f'{label} must be a UUID.')


def bounded_int(value, default, minimum=1, maximum=200):
    try: result = int(value) if value is not None else default
    except (ValueError, TypeError): raise APIError('Pagination values must be integers.')
    if not minimum <= result <= maximum: raise APIError(f'Value must be between {minimum} and {maximum}.')
    return result


def permitted_tournaments():
    value = getattr(getattr(g, 'membership', None), 'tournament_ids', None)
    return list(value or [])


def scoped(model, value):
    value = ident(value)
    row = get_db().scalar(select(model).where(model.id == value, model.organization_id == g.organization.id))
    if row is None: raise APIError('Record is unavailable in this organization.', 404, 'not_found')
    tournament_id = row.id if isinstance(row, Tournament) else getattr(row, 'tournament_id', None)
    if tournament_id and permitted_tournaments() and tournament_id not in permitted_tournaments():
        raise APIError('Tournament is outside your assigned scope.', 403, 'forbidden')
    return row


def clean_filters(values):
    result = {}
    for key in FILTERS:
        value = values.get(key)
        if value not in (None, ''):
            if not isinstance(value, (str, int)) or len(str(value)) > 300:
                raise APIError(f'Invalid report filter: {key}.')
            result[key] = str(value)
    for key, model in [('tournament_id',Tournament),('division_id',Division),('stage_id',Stage),
                       ('venue_id',Venue),('court_id',Court)]:
        if key in result:
            result[key]=ident(result[key],key);scoped(model,result[key])
    if result.get('team_id'):
        team_id = ident(result['team_id'], 'team_id')
        entry = get_db().scalar(select(Entry).where(Entry.id == team_id, Entry.organization_id == g.organization.id))
        if entry: scoped(Entry, entry.id)
        elif not get_db().scalar(select(Team.id).where(Team.id == team_id, Team.organization_id == g.organization.id)):
            raise APIError('Team is unavailable.', 404, 'not_found')
        result['team_id'] = team_id
    # Reject conflicting references rather than silently returning an unrelated scope.
    if result.get('division_id'):
        division = scoped(Division, result['division_id'])
        if result.get('tournament_id') and division.tournament_id != result['tournament_id']:
            raise APIError('Division does not belong to the selected tournament.')
    if result.get('stage_id'):
        stage = scoped(Stage, result['stage_id'])
        if result.get('tournament_id') and stage.tournament_id != result['tournament_id']:
            raise APIError('Stage does not belong to the selected tournament.')
        if result.get('division_id') and stage.division_id != result['division_id']:
            raise APIError('Stage does not belong to the selected division.')
    return result


def date_bound(value, tz, end=False):
    if not value: return None
    try:
        if len(value) == 10:
            day = date.fromisoformat(value)
            if end: day += timedelta(days=1)
            return datetime.combine(day, time.min, ZoneInfo(tz)).astimezone(timezone.utc)
        stamp = datetime.fromisoformat(value.replace('Z', '+00:00'))
        if stamp.tzinfo is None: stamp = stamp.replace(tzinfo=ZoneInfo(tz))
        return stamp.astimezone(timezone.utc)
    except (ValueError, TypeError): raise APIError('Dates must be ISO dates or timestamps.')


def report_source(filters):
    db = get_db()
    query = select(Tournament).where(Tournament.organization_id == g.organization.id)
    if permitted_tournaments(): query = query.where(Tournament.id.in_(permitted_tournaments()))
    if filters.get('tournament_id'): query = query.where(Tournament.id == filters['tournament_id'])
    tournaments = list(db.scalars(query))
    tids = [t.id for t in tournaments]
    query = select(Match).where(Match.organization_id == g.organization.id, Match.tournament_id.in_(tids))
    for key in ('division_id','stage_id','status','result_status','outcome'):
        if filters.get(key): query = query.where(getattr(Match,key) == filters[key])
    entries = {row.id: row for row in db.scalars(select(Entry).where(Entry.organization_id == g.organization.id, Entry.tournament_id.in_(tids)))}
    if filters.get('team_id'):
        eids = [e.id for e in entries.values() if e.id == filters['team_id'] or e.team_id == filters['team_id']]
        query = query.where(or_(Match.entry_a_id.in_(eids), Match.entry_b_id.in_(eids)))
    tz = tournaments[0].timezone if len(tournaments) == 1 else g.organization.timezone
    begin, finish = date_bound(filters.get('from'),tz), date_bound(filters.get('to'),tz,True)
    if begin and finish and begin >= finish: raise APIError('Start date must precede the end date.')
    if begin: query = query.where(Match.planned_at >= begin)
    if finish: query = query.where(Match.planned_at < finish)
    assignments = list(db.scalars(select(CourtAssignment).where(CourtAssignment.organization_id == g.organization.id)))
    courts = {c.id:c for c in db.scalars(select(Court).where(Court.organization_id == g.organization.id))}
    assignment_by_match = {}
    for a in sorted(assignments,key=lambda a:a.created_at): assignment_by_match[a.match_id] = a
    if filters.get('court_id') or filters.get('venue_id'):
        mids = [a.match_id for a in assignments if
                (not filters.get('court_id') or a.court_id == filters['court_id']) and
                (not filters.get('venue_id') or courts[a.court_id].venue_id == filters['venue_id'])]
        query = query.where(Match.id.in_(mids))
    matches = list(db.scalars(query.order_by(Match.planned_at,Match.id).limit(MAX_REPORT_ROWS+1)))
    if len(matches) > MAX_REPORT_ROWS: raise APIError('Report exceeds 50,000 matches; narrow the date or tournament filters.',413,'report_limit')
    stages = {s.id:s for s in db.scalars(select(Stage).where(Stage.organization_id == g.organization.id, Stage.tournament_id.in_(tids)))}
    divisions = {d.id:d for d in db.scalars(select(Division).where(Division.organization_id == g.organization.id, Division.tournament_id.in_(tids)))}
    return tournaments, matches, entries, divisions, stages, courts, assignment_by_match, assignments, tz


def build_report(report_type, filters, public=False):
    """Read all rows and source generations from one PostgreSQL snapshot."""
    original=get_db()
    snapshot=Session(current_app.extensions['db_engine'].execution_options(isolation_level='REPEATABLE READ'),expire_on_commit=False)
    snapshot.info['organization_id']=g.organization.id
    g.db=snapshot
    try: return _build_report(report_type,filters,public)
    finally:
        snapshot.rollback();snapshot.close();g.db=original


def _build_report(report_type, filters, public=False):
    if report_type not in TYPES or (public and report_type not in PUBLIC_TYPES): raise APIError('Unsupported report type.')
    if report_type=='standings' and (any(filters.get(key) for key in ('from','to','status','outcome','court_id','venue_id')) or filters.get('result_status') not in {None,'','finalized'}):
        raise APIError('Official standings use complete finalized division/stage results. Filter by tournament, division, stage or team; use a match report for date, court, lifecycle or outcome filters.')
    if public:
        db=get_db();tid=ident(filters.get('tournament_id'))
        for key,model in [('division_id',Division),('stage_id',Stage)]:
            if filters.get(key) and not db.scalar(select(model.id).where(model.id==ident(filters[key]),model.tournament_id==tid)):
                raise APIError('Filter is unavailable in this public event.',404,'not_found')
        if filters.get('team_id'):
            value=ident(filters['team_id'])
            if not db.scalar(select(Entry.id).where(Entry.tournament_id==tid,or_(Entry.id==value,Entry.team_id==value))):
                raise APIError('Filter is unavailable in this public event.',404,'not_found')
        tournament=db.get(Tournament,tid)
        venue_ids=(tournament.settings or {}).get('venue_ids',[])
        used_courts=select(CourtAssignment.court_id).join(Match,Match.id==CourtAssignment.match_id).where(Match.tournament_id==tid)
        allowed_courts=list(db.scalars(select(Court).where(or_(Court.venue_id.in_(venue_ids),Court.id.in_(used_courts)))))
        if filters.get('court_id') and ident(filters['court_id']) not in {c.id for c in allowed_courts}:
            raise APIError('Filter is unavailable in this public event.',404,'not_found')
        if filters.get('venue_id') and ident(filters['venue_id']) not in set(venue_ids)|{c.venue_id for c in allowed_courts}:
            raise APIError('Filter is unavailable in this public event.',404,'not_found')
    filters = clean_filters(filters)
    ts, matches, entries, divisions, stages, courts, current_assignments, assignments, tz = report_source(filters)
    tournament_by_id = {t.id:t for t in ts}
    rows = []
    columns = []
    ranking_policies={}
    if report_type in {'matches','schedule','referees','outstanding_results'}:
        columns = ['match_id','match_version','tournament','division','stage','team_a','team_b','court','planned_at','actual_start','actual_end',
                   'planned_minutes','actual_minutes','delay_minutes','status','result_status','sets','winner','outcome']
        if report_type == 'matches': columns += ['rules_used','result_revision']
        if not public: columns += ['referee_name']
        for m in matches:
            if report_type == 'outstanding_results' and m.result_status not in {'submitted','pending'}: continue
            a = current_assignments.get(m.id)
            elapsed = (m.actual_end-m.actual_start).total_seconds()/60 if m.actual_end and m.actual_start else None
            row = {'match_id':m.id,'match_version':m.version,'tournament':tournament_by_id[m.tournament_id].name,'division':divisions[m.division_id].name,
                   'stage':stages[m.stage_id].name,'team_a':entries[m.entry_a_id].name if m.entry_a_id else 'TBD',
                   'team_b':entries[m.entry_b_id].name if m.entry_b_id else 'TBD',
                   'court':courts[a.court_id].name if a else 'Court to be assigned','planned_at':iso(m.planned_at),
                   'actual_start':iso(m.actual_start),'actual_end':iso(m.actual_end),'planned_minutes':m.planned_duration,
                   'actual_minutes':round(elapsed,2) if elapsed is not None else None,
                   'delay_minutes':round(max(0,(m.actual_start-m.planned_at).total_seconds()/60),2) if m.actual_start and m.planned_at else None,
                   'status':m.status,'result_status':m.result_status,'sets':(m.state or {}).get('sets',[]),
                   'winner':entries[m.winner_id].name if m.winner_id in entries else None,'outcome':m.outcome}
            if report_type == 'matches': row.update(rules_used=m.rule_snapshot or m.rules,result_revision=m.result_revision)
            if not public: row['referee_name']=m.referee_name
            rows.append(row)
    elif report_type in {'summary','history'}:
        columns = ['tournament_id','tournament','competition_type','start_date','end_date','divisions','teams','games','live','ended','finalized','completion_percent','outcomes']
        if report_type=='history': columns+=['average_actual_minutes','finalized_sets','finalized_rally_points']
        for t in ts:
            selected = [m for m in matches if m.tournament_id == t.id]
            ds = [d for d in divisions.values() if d.tournament_id == t.id and (not filters.get('division_id') or d.id == filters['division_id'])]
            es = [e for e in entries.values() if e.division_id in {d.id for d in ds}]
            if any(key in filters for key in ('stage_id','team_id','from','to','court_id','venue_id','status','result_status','outcome')):
                participants={eid for m in selected for eid in (m.entry_a_id,m.entry_b_id) if eid}
                es=[e for e in es if e.id in participants]
                ds=[d for d in ds if any(e.division_id==d.id for e in es)]
            outcomes = {}
            for m in selected:
                if m.outcome: outcomes[m.outcome] = outcomes.get(m.outcome,0)+1
            ended = sum(m.status in {'ended','cancelled','abandoned'} for m in selected)
            summary={'tournament_id':t.id,'tournament':t.name,'competition_type':t.competition_type,'start_date':iso(t.start_date),
                         'end_date':iso(t.end_date),'divisions':len(ds),'teams':len(es),'games':len(selected),
                         'live':sum(m.status in {'in_progress','paused'} for m in selected),'ended':ended,
                         'finalized':sum(m.result_status=='finalized' for m in selected),
                         'completion_percent':round(100*ended/len(selected),2) if selected else 0,'outcomes':outcomes}
            if report_type=='history':
                durations=[(m.actual_end-m.actual_start).total_seconds()/60 for m in selected if m.actual_start and m.actual_end]
                sets=[s for m in selected if m.result_status=='finalized' and m.outcome!='forfeit' for s in (m.state or {}).get('sets',[]) if s.get('complete')]
                summary.update(average_actual_minutes=round(sum(durations)/len(durations),2) if durations else None,
                               finalized_sets=len(sets),finalized_rally_points=sum(s.get('a',0)+s.get('b',0) for s in sets))
            rows.append(summary)
    elif report_type == 'court_use':
        columns = ['assignment_id','match_id','court','venue_id','status','reserved_at','expires_at','started_at','released_at','occupied_minutes']
        mids = {m.id for m in matches}
        for a in assignments:
            if a.match_id not in mids: continue
            c = courts[a.court_id]
            if filters.get('court_id') and c.id!=filters['court_id']: continue
            if filters.get('venue_id') and c.venue_id!=filters['venue_id']: continue
            rows.append({'assignment_id':a.id,'match_id':a.match_id,'court':c.name,'venue_id':c.venue_id,'status':a.status,
                         'reserved_at':iso(a.created_at),'expires_at':iso(a.expires_at),'started_at':iso(a.started_at),
                         'released_at':iso(a.released_at),'occupied_minutes':round((a.released_at-a.started_at).total_seconds()/60,2)
                         if a.started_at and a.released_at else None})
    elif report_type == 'audit':
        if public: raise APIError('Audit reports are private.',403,'forbidden')
        columns = ['audit_id','at','actor_type','actor_id','session_id','action','entity_type','entity_id','reason','changes']
        query = select(Audit).where(Audit.organization_id == g.organization.id)
        allowed_matches = {m.id for m in matches}
        allowed_divisions = {d.id for d in divisions.values() if (not filters.get('division_id') or d.id == filters['division_id'])}
        allowed_stages = {s.id for s in stages.values() if s.division_id in allowed_divisions and (not filters.get('stage_id') or s.id == filters['stage_id'])}
        competition_filter = any(filters.get(key) for key in
            ('tournament_id','division_id','stage_id','team_id','venue_id','court_id','status','result_status','outcome'))
        if permitted_tournaments() or competition_filter:
            scope_conditions = [(Audit.entity_type=='match') & Audit.entity_id.in_(allowed_matches),
                (Audit.entity_type=='tournament') & Audit.entity_id.in_(tournament_by_id),
                (Audit.entity_type=='division') & Audit.entity_id.in_(allowed_divisions),
                (Audit.entity_type=='stage') & Audit.entity_id.in_(allowed_stages)]
            if not competition_filter:
                scope_conditions.append(Audit.entity_type=='support_case')
            query = query.where(or_(*scope_conditions))
        begin,finish=date_bound(filters.get('from'),tz),date_bound(filters.get('to'),tz,True)
        if begin: query=query.where(Audit.created_at>=begin)
        if finish: query=query.where(Audit.created_at<finish)
        for a in get_db().scalars(query.order_by(Audit.created_at.desc()).limit(MAX_REPORT_ROWS)):
            rows.append({'audit_id':a.id,'at':iso(a.created_at),'actor_type':a.actor_type,'actor_id':a.actor_id,'session_id':a.session_id,
                         'action':a.action,'entity_type':a.entity_type,'entity_id':a.entity_id,'reason':a.reason,
                         'changes':{'before':without_secrets(a.before),'after':without_secrets(a.after)}})
    elif report_type == 'team_performance':
        from .services.rules import effective_rules
        columns=['entry_id','team','tournament','division','played','wins','losses','draws','sets_won','sets_lost','points_scored','points_conceded']
        stats={}
        for m in matches:
            if m.result_status!='finalized' or not m.entry_a_id or not m.entry_b_id or m.outcome in {'bye','abandoned','cancelled'}: continue
            policy=effective_rules(tournament_by_id[m.tournament_id],divisions[m.division_id],stages[m.stage_id] if filters.get('stage_id') else None)
            for side,eid,opposite in [('a',m.entry_a_id,'b'),('b',m.entry_b_id,'a')]:
                if not eid: continue
                e=entries[eid]
                row=stats.setdefault(eid,dict.fromkeys(columns,0))
                row.update(entry_id=eid,team=e.name,tournament=tournament_by_id[e.tournament_id].name,division=divisions[e.division_id].name)
                row['played']+=1
                row['wins' if m.winner_id==eid else 'draws' if not m.winner_id else 'losses']+=1
                for s in (m.state or {}).get('sets',[]):
                    if not s.get('complete') or (m.outcome=='forfeit' and not policy['include_forfeit_scores']):continue
                    row['points_scored']+=s.get(side,0); row['points_conceded']+=s.get(opposite,0)
                    if s.get('complete'):
                        row['sets_won' if s.get('winner')==side else 'sets_lost']+=1
        rows=list(stats.values())
        if filters.get('team_id'): rows=[r for r in rows if r['entry_id']==filters['team_id'] or entries[r['entry_id']].team_id==filters['team_id']]
    elif report_type == 'standings':
        # Reuse the authoritative ranking service. Scope/data filters never recompute
        # an invented partial official table; rows are selected from the full table.
        from .services.rankings import standings
        columns=['division','rank','entry_id','name','played','won','lost','drawn','sets_won','sets_lost','points_for','points_against','standings_points','set_difference','point_difference','set_quotient','point_quotient','win_percentage','tied','explanation']
        for division in divisions.values():
            if filters.get('division_id') and division.id!=filters['division_id']: continue
            if filters.get('stage_id') and stages[filters['stage_id']].division_id!=division.id: continue
            ranking=standings(division,stage_id=filters.get('stage_id'))
            ranking_policies[division.id]={'policy':ranking['policy'],'generation':ranking['generation'],'explanation':ranking['explanation']}
            ranking_rows=ranking.get('rows',ranking.get('items',[])) if isinstance(ranking,dict) else ranking
            for source in ranking_rows:
                eid=source.get('entry_id',source.get('id'))
                if filters.get('team_id') and eid!=filters['team_id'] and (not entries.get(eid) or entries[eid].team_id!=filters['team_id']): continue
                row={key:source.get(key) for key in columns}; row.update(division=division.name)
                row['explanation']=source.get('explanation',ranking.get('explanation') if isinstance(ranking,dict) else None)
                rows.append(row)
    if filters.get('search'):
        needle=filters['search'].casefold()
        rows=[r for r in rows if needle in json.dumps(r,ensure_ascii=False,default=str).casefold()]
    metadata={'type':report_type,'organization':{'id':g.organization.id,'name':g.organization.name,'demo':g.organization.demo},
              'tournaments':[{'id':t.id,'name':t.name} for t in ts],'generated_at':iso(now()),'timezone':tz,'filters':filters,
              'official_status':'official finalized results' if report_type in {'standings','team_performance'} else
                                'filtered lifecycle totals; finalized fields are official' if report_type in {'summary','history'} else
                                'official and provisional; see result_status',
              'source_generations':{t.id:t.standings_generation for t in ts},
              'source_match_versions':{m.id:{'version':m.version,'result_revision':m.result_revision} for m in matches},
              'public_safe':public,'credit':'Designed and developed by Vamsi Pinnaka',
              'ranking_policies':ranking_policies,
              'metric_definitions':{'completion_percent':'Ended, cancelled or abandoned games divided by filtered games.',
                                    'delay_minutes':'Nonnegative actual start minus planned start.',
                                    'occupied_minutes':'Released minus started; open assignments have no finalized duration.'}}
    total_columns={'games','live','ended','finalized','played','won','lost','drawn','wins','losses','draws','sets_won','sets_lost','points_for','points_against','points_scored','points_conceded','standings_points','finalized_sets','finalized_rally_points'}
    return {'columns':columns,'rows':rows,'total':len(rows),'metadata':metadata,
            'totals':{c:sum(r.get(c) or 0 for r in rows) for c in columns if c in total_columns}}


def present_report(report, values, paginate=True):
    cols=values.get('columns')
    if cols:
        cols=[x.strip() for x in cols.split(',')] if isinstance(cols,str) else cols
        if not isinstance(cols,list) or not cols or any(c not in report['columns'] for c in cols): raise APIError('Choose known report columns.')
        report['columns']=list(dict.fromkeys(cols))
    sort=values.get('sort')
    if sort:
        if sort not in report['columns']: raise APIError('Unknown sort column.')
        direction=values.get('direction','asc')
        if direction not in {'asc','desc'}: raise APIError('Sort direction must be asc or desc.')
        def sort_key(row):
            value=row.get(sort)
            return (value is None,0 if isinstance(value,(int,float)) else 1,value if isinstance(value,(int,float)) else str(value or '').casefold())
        report['rows'].sort(key=sort_key,reverse=direction=='desc')
    report['rows']=[{c:r.get(c) for c in report['columns']} for r in report['rows']]
    if paginate:
        page=bounded_int(values.get('page'),1,maximum=100000)
        size=bounded_int(values.get('page_size'),50)
        report['rows']=report['rows'][(page-1)*size:page*size]
        report.update(page=page,page_size=size)
    return report


def spreadsheet_safe(value):
    if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False,sort_keys=True)
    if value is None: return ''
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r','\n')): return "'"+value
    return value


def pdf_column_groups(columns,report_type):
    """Keep a selected stable identifier on every wide-report page group."""
    preferred={'matches':'match_id','schedule':'match_id','referees':'match_id','outstanding_results':'match_id',
        'summary':'tournament_id','history':'tournament_id','court_use':'assignment_id','audit':'audit_id',
        'standings':'entry_id','team_performance':'entry_id'}.get(report_type)
    identity=preferred if preferred in columns else next((key for key in
        ('assignment_id','audit_id','entry_id','match_id','tournament_id') if key in columns),None)
    # If a saved column selection omits identifiers, an explicit source-row
    # ordinal still links the grouped tables without adding excluded data.
    identity=identity or '__report_row'
    groups=[]
    for offset in range(0,len(columns),6):
        selected=columns[offset:offset+6]
        if offset and selected==[identity]:continue
        groups.append(selected if identity in selected else [identity,*selected])
    return groups


def pdf_text(value):
    # Noto Sans supplies text glyphs. The PDF standard Symbol font supplies
    # the ranking infinity symbol without changing user text or metric values.
    return escape(str(value)).replace('∞','<font name="Symbol">∞</font>')


def render_report(report, format):
    if format not in FORMATS: raise APIError('Choose csv, xlsx or pdf.')
    if format=='csv':
        output=StringIO(newline=''); writer=csv.writer(output)
        writer.writerow(report['columns'])
        for row in report['rows']: writer.writerow([spreadsheet_safe(row.get(c)) for c in report['columns']])
        # CSV has no hidden metadata layer; append an explicit non-data section.
        writer.writerow([]);writer.writerow(['REPORT METADATA'])
        for key,value in report['metadata'].items(): writer.writerow([key,spreadsheet_safe(value)])
        return output.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb=Workbook(); sheet=wb.active;sheet.title='Report';sheet.append(report['columns'])
        for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='12604B')
        for row in report['rows']: sheet.append([spreadsheet_safe(row.get(c)) for c in report['columns']])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=24
        meta=wb.create_sheet('Metadata');meta.append(['Field','Value'])
        for key,value in report['metadata'].items():
            value=spreadsheet_safe(value)
            if isinstance(value,str) and len(value)>30000:
                count=(len(value)+29999)//30000
                for part in range(count):
                    meta.append([f'{key} [part {part+1}/{count}]',value[part*30000:(part+1)*30000]])
                    # Explicit string cells preserve all metadata without executable
                    # formula inference or Excel's 32,767-character truncation.
                    meta.cell(meta.max_row,2).data_type='s'
            else:meta.append([key,value])
        output=BytesIO();wb.save(output)
        return output.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    font_dir=Path(__file__).with_name('fonts')
    if 'VolleyballNoto' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('VolleyballNoto',str(font_dir/'NotoSans-Regular.ttf')))
        pdfmetrics.registerFont(TTFont('VolleyballNotoBold',str(font_dir/'NotoSans-Bold.ttf')))
    output=BytesIO();styles=getSampleStyleSheet()
    styles['BodyText'].fontName='VolleyballNoto';styles['BodyText'].fontSize=9;styles['BodyText'].leading=12
    styles['Heading1'].fontName='VolleyballNotoBold';styles['Heading2'].fontName='VolleyballNotoBold'
    doc=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=32)
    items=[Paragraph('Volleyball Manager - '+pdf_text(report['metadata']['type'].replace('_',' ').title()),styles['Heading1'])]
    for key in ('organization','generated_at','timezone','filters','official_status','source_generations'):
        value=report['metadata'].get(key)
        items.append(Paragraph(pdf_text(key.replace('_',' ').title()+': '+str(value)),styles['BodyText']))
    if report['metadata'].get('ranking_policies'):
        items.append(Paragraph('Ranking policies and explanations',styles['Heading2']))
        for division_id,policy in report['metadata']['ranking_policies'].items():
            items.append(Paragraph(pdf_text(division_id+': '+json.dumps(policy,ensure_ascii=False,sort_keys=True)),styles['BodyText']))
    items.append(Spacer(1,12))
    # Split wide reports into grouped column tables while retaining a row identity.
    for group,selected in enumerate(pdf_column_groups(report['columns'],report['metadata']['type'])):
        if group: items.extend([PageBreak(),Paragraph('Additional columns',styles['Heading2'])])
        data=[[Paragraph('Row' if c=='__report_row' else pdf_text(c.replace('_',' ')),styles['BodyText']) for c in selected]]
        for index,row in enumerate(report['rows'],1):
            data.append([Paragraph(pdf_text(index if c=='__report_row' else spreadsheet_safe(row.get(c))),styles['BodyText']) for c in selected])
        if len(data)==1: data.append([Paragraph('No data' if i==0 else '',styles['BodyText']) for i in range(len(selected))])
        table=Table(data,repeatRows=1,colWidths=[(landscape(A4)[0]-48)/len(selected)]*len(selected))
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#def5ec')),('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        items.append(table)
    def footer(canvas,doc):
        canvas.setFont('VolleyballNoto',8);canvas.drawString(24,18,'Designed and developed by Vamsi Pinnaka');canvas.drawRightString(landscape(A4)[0]-24,18,str(doc.page))
    doc.build(items,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue(),'application/pdf'


@bp.get('/reports')
@require_org('read')
def reports():
    return present_report(build_report(request.args.get('type','matches'),request.args),request.args)


@bp.get('/reports/export')
@require_org('read')
def export_direct():
    from .auth import _throttle
    _throttle('private_report_export',g.user.id,10,minutes=1)
    report=present_report(build_report(request.args.get('type','matches'),request.args),request.args,False)
    if report['total']>2000: raise APIError('Use a queued export for more than 2,000 rows.',413,'use_export_job')
    format=request.args.get('format','csv');payload,mime=render_report(report,format)
    audit('report.export','report',None,after={'type':report['metadata']['type'],'filters':report['metadata']['filters'],'rows':report['total'],'format':format})
    get_db().commit()
    return send_file(BytesIO(payload),mimetype=mime,as_attachment=True,download_name='volleyball-report.'+format)


@bp.get('/public/<slug>/reports/export')
def public_export(slug):
    from .competition import public_tournament
    tournament=public_tournament(slug)
    report_type=request.args.get('type','schedule')
    if report_type not in PUBLIC_TYPES or report_type not in (tournament.display or {}).get('public_reports',[]):
        raise APIError('This public report is not enabled.',404,'not_found')
    from .auth import _throttle
    _throttle('public_report_export',request.remote_addr or 'unknown',10,minutes=1)
    get_db().commit()
    g.organization=get_db().get(Organization,tournament.organization_id)
    g.membership=None
    filters=dict(request.args);filters['tournament_id']=tournament.id
    report=present_report(build_report(report_type,filters,True),request.args,False)
    if report['total']>2000: raise APIError('Narrow the public report filters.',413,'report_limit')
    format=request.args.get('format','csv');payload,mime=render_report(report,format)
    get_db().refresh(tournament)
    get_db().refresh(g.organization)
    if tournament.visibility!='public' or tournament.status=='draft' or not g.organization.active or report_type not in (tournament.display or {}).get('public_reports',[]):
        raise APIError('Public report is no longer available.',404,'not_found')
    return send_file(BytesIO(payload),mimetype=mime,as_attachment=True,download_name=tournament.slug+'-'+report_type+'.'+format)


def job_json(job):
    value={key:getattr(job,key) for key in ('id','kind','status','type','format','filters','progress','error')}
    value.update(created_at=iso(job.created_at),completed_at=iso(job.completed_at),expires_at=iso(job.expires_at))
    if job.kind=='import': value.update(preview=job.data.get('preview',[]),errors=job.data.get('errors',[]),row_count=len(job.data.get('rows',[])))
    else: value['metadata']=job.data.get('metadata')
    return value


def own_job(job_id,kind=None):
    job=scoped(Job,job_id)
    if job.user_id!=g.user.id or (kind and job.kind!=kind): raise APIError('Job is unavailable.',404,'not_found')
    if job.kind=='export': check_job_scope(job)
    if job.kind=='import':
        import_permission(job.type)
        if permitted_tournaments() and any(row.get('tournament_id') and row['tournament_id'] not in permitted_tournaments() for row in job.data.get('preview',[])):
            raise APIError('Import preview is outside your current tournament assignment.',403,'forbidden')
    return job


@bp.post('/exports')
@require_org('read')
def create_export():
    values=body();report_type=values.get('type','matches');format=values.get('format','csv')
    if report_type not in TYPES or format not in FORMATS: raise APIError('Choose a supported report type and format.')
    if not isinstance(values.get('filters',{}),dict): raise APIError('Filters must be an object.')
    filters=clean_filters(values.get('filters',{}))
    if get_db().scalar(select(func.count()).select_from(Job).where(Job.organization_id==g.organization.id,Job.user_id==g.user.id,Job.kind=='export',Job.status.in_(['queued','running'])))>=5:
        raise APIError('Five exports are already pending; wait or cancel one.',429,'rate_limit')
    layout={key:values[key] for key in ('sort','direction','columns') if key in values}
    job=Job(organization_id=g.organization.id,user_id=g.user.id,kind='export',type=report_type,format=format,filters=filters,
            data={'view':layout,'auth_session_id':g.auth_session.id},expires_at=now()+timedelta(hours=24))
    get_db().add(job);get_db().flush();audit('export.queued','job',job.id,after={'type':report_type,'format':format,'filters':filters});get_db().commit()
    return job_json(job),202


@bp.get('/exports')
@require_org('read')
def exports():
    jobs=list(get_db().scalars(select(Job).where(Job.organization_id==g.organization.id,Job.user_id==g.user.id,Job.kind=='export').order_by(Job.created_at.desc()).limit(100)))
    visible=[]
    for job in jobs:
        try: check_job_scope(job)
        except APIError: continue
        visible.append(job_json(job))
    return {'items':visible,'total':len(visible)}


@bp.get('/exports/<job_id>')
@require_org('read')
def get_export(job_id): return job_json(own_job(job_id,'export'))


@bp.delete('/exports/<job_id>')
@require_org('read')
def cancel_export(job_id):
    job=own_job(job_id,'export')
    if job.status not in {'queued','running'}: raise APIError('Only queued or running exports can be cancelled.',409,'conflict')
    job.status='cancelled';audit('export.cancelled','job',job.id);get_db().commit();return job_json(job)


def private_root():
    root=Path(current_app.config.get('PRIVATE_STORAGE_PATH',Path(current_app.instance_path)/'private')).resolve()
    root.mkdir(parents=True,exist_ok=True,mode=0o700)
    return root


def safe_file(job):
    if not job.file_path: raise APIError('Export file is unavailable.',404,'not_found')
    root=private_root(); path=(root/job.file_path).resolve()
    if not path.is_relative_to(root) or path.name != job.id+'.'+job.format: raise APIError('Invalid export file reference.',404,'not_found')
    return path


def check_job_scope(job):
    clean_filters(job.filters)
    generated=set((job.data.get('metadata') or {}).get('source_generations',{}))
    if permitted_tournaments() and not generated.issubset(set(permitted_tournaments())):
        raise APIError('Your current tournament assignment does not permit this export.',403,'forbidden')


@bp.get('/exports/<job_id>/download')
@require_org('read')
def download_export(job_id):
    job=own_job(job_id,'export');check_job_scope(job)
    if not job.expires_at or job.expires_at<=now(): raise APIError('Export has expired; generate a new report.',410,'expired')
    if job.status!='completed': raise APIError('Export is not completed.',409,'not_ready')
    path=safe_file(job)
    if not path.is_file(): raise APIError('File is missing; generate a new export.',410,'missing_file')
    audit('export.downloaded','job',job.id);get_db().commit()
    return send_file(path,as_attachment=True,download_name='volleyball-'+job.type+'.'+job.format)


def authorize_worker(job):
    db=get_db();set_scope(job.organization_id)
    org=db.scalar(select(Organization).where(Organization.id==job.organization_id).execution_options(populate_existing=True))
    user=db.scalar(select(User).where(User.id==job.user_id).execution_options(populate_existing=True))
    member=db.scalar(select(Membership).where(Membership.organization_id==job.organization_id,Membership.user_id==job.user_id,Membership.status=='active').execution_options(populate_existing=True))
    if not org or not org.active or not user or not user.active or not member:
        raise APIError('Requester access has been revoked.',403,'forbidden')
    auth_session=db.scalar(select(AuthSession).where(AuthSession.id==job.data['auth_session_id']).execution_options(populate_existing=True)) if job.data.get('auth_session_id') else None
    if not auth_session or auth_session.user_id!=user.id or auth_session.revoked_at or auth_session.expires_at<=now():
        raise APIError('Requester session is no longer valid; sign in and create a new export.',403,'forbidden')
    if org.settings.get('require_mfa') and (not user.mfa_enabled or not auth_session.mfa_verified):
        raise APIError('Requester must enable administrator MFA.',403,'mfa_required')
    g.user=user;g.organization=org;g.membership=member
    clean_filters(job.filters)


def process_export_jobs(limit=10):
    """Process jobs without leaking the last worker organization into its caller."""
    db=get_db();scope=db.info.get('organization_id','')
    context={key:getattr(g,key,None) for key in ('user','organization','membership','auth_session')}
    try: return _process_export_jobs(limit)
    finally:
        db.rollback();set_scope(scope or '')
        for key,value in context.items():
            if value is None: g.pop(key,None)
            else: setattr(g,key,value)


def _process_export_jobs(limit=10):
    """Worker entry point; callers must run inside an application context."""
    db=get_db();processed=0
    orgs=list(db.scalars(select(Organization)))
    for org in orgs:
        set_scope(org.id)
        stale=list(db.scalars(select(Job).where(Job.organization_id==org.id,Job.kind=='export',Job.status=='running',Job.updated_at<now()-timedelta(minutes=15)).with_for_update(skip_locked=True)))
        for job in stale: job.status='queued';job.progress=0
        db.commit()
        while processed<limit:
            job=db.scalar(select(Job).where(Job.organization_id==org.id,Job.kind=='export',Job.status=='queued').order_by(Job.created_at).with_for_update(skip_locked=True).limit(1))
            if not job: db.rollback();break
            job_id=job.id;lease_id=str(uuid.uuid4());job.status='running';job.progress=5
            job.data={**job.data,'lease_id':lease_id};db.commit()
            output=None;temporary=None
            try:
                authorize_worker(job)
                if job.expires_at and job.expires_at<=now(): raise APIError('Export expired before processing.',410,'expired')
                report=present_report(build_report(job.type,job.filters),job.data.get('view',{}),False)
                job=db.scalar(select(Job).where(Job.id==job_id).execution_options(populate_existing=True).with_for_update())
                if job.status!='running' or job.data.get('lease_id')!=lease_id:
                    db.commit();processed+=1;continue
                job.progress=80;db.commit()
                payload,mime=render_report(report,job.format)
                root=private_root();directory=root/'exports'/job.organization_id;directory.mkdir(parents=True,exist_ok=True,mode=0o700)
                output=directory/(job.id+'.'+job.format);temporary=directory/(job.id+'.'+job.format+'.'+lease_id+'.partial')
                temporary.write_bytes(payload);temporary.chmod(0o600)
                job=db.scalar(select(Job).where(Job.id==job_id).execution_options(populate_existing=True).with_for_update())
                authorize_worker(job)
                if job.status!='running' or job.data.get('lease_id')!=lease_id:
                    temporary.unlink(missing_ok=True);db.commit();processed+=1;continue
                temporary.replace(output);temporary=None
                job.file_path=output.relative_to(root).as_posix();job.data={**job.data,'metadata':report['metadata'],'mime':mime}
                job.status='completed';job.progress=100;job.completed_at=now();job.error=None
                audit('export.completed','job',job.id,after={'rows':report['total'],'format':job.format});db.commit()
            except Exception as error:
                db.rollback();set_scope(org.id)
                job=db.scalar(select(Job).where(Job.id==job_id).with_for_update())
                if job and job.status=='running' and job.data.get('lease_id')==lease_id:
                    job.status='failed';job.error=error.message if isinstance(error,APIError) else 'Export generation failed; review operator logs and retry.';job.completed_at=now()
                    audit('export.failed','job',job.id,after={'error':job.error});db.commit()
                if temporary: temporary.unlink(missing_ok=True)
                if output and job and job.status=='failed' and job.data.get('lease_id')==lease_id: output.unlink(missing_ok=True)
                if not isinstance(error,APIError): current_app.logger.error('export_generation_failed job_id=%s organization_id=%s error_type=%s',job_id,org.id,type(error).__name__)
            processed+=1
        expired=list(db.scalars(select(Job).where(Job.organization_id==org.id,Job.kind=='export',Job.file_path.is_not(None),Job.expires_at<now()).with_for_update(skip_locked=True)))
        for job in expired:
            safe_file(job).unlink(missing_ok=True);job.file_path=None;job.status='expired'
        db.commit()
    return processed


def init_reports(app):
    @app.cli.command('worker')
    @click.option('--once',is_flag=True,help='Process one batch and exit.')
    def worker(once):
        import time as clock
        while True:
            count=process_export_jobs();click.echo(f'Processed {count} export jobs.')
            if once: return
            clock.sleep(2)


def import_permission(kind):
    roles={'teams':{'admin','manager'},'entries':{'admin','manager'},'fixtures':{'admin','manager','scheduler'},'referees':{'admin','manager','scheduler'}}
    if kind not in roles: raise APIError('Unsupported import type.')
    if g.membership.role not in roles[kind]: raise APIError('Your role cannot perform this import.',403,'forbidden')


def validate_import(kind,rows):
    import_permission(kind)
    errors=[];validated=[];seen=set();db=get_db()
    for number,row in enumerate(rows,2):
        try:
            if set(row)-set(TEMPLATES[kind]): raise APIError('Unexpected columns.')
            data={key:(value or '').strip() for key,value in row.items()}
            if any(len(v)>500 for v in data.values()): raise APIError('Field exceeds 500 characters.')
            if kind=='teams':
                if not data.get('name') or len(data['name'])>160: raise APIError('A team name of at most 160 characters is required.')
                if len(data.get('abbreviation',''))>20: raise APIError('Abbreviation exceeds 20 characters.')
                color=data.get('color') or '#11b981'
                if len(color)!=7 or color[0]!='#' or any(c not in '0123456789abcdefABCDEF' for c in color[1:]): raise APIError('Color must be #RRGGBB.')
                data['color']=color;key=data['name'].casefold()
                if db.scalar(select(Team.id).where(Team.organization_id==g.organization.id,func.lower(Team.name)==data['name'].lower())): raise APIError('Team name already exists; imports do not overwrite records.')
            elif kind=='entries':
                division=scoped(Division,data.get('division_id'));team=scoped(Team,data.get('team_id'))
                if team.archived: raise APIError('Archived team cannot be entered.')
                data['seed']=bounded_int(data.get('seed') or 1,1,1,100000)
                data['name']=data.get('name') or team.name
                if len(data['name'])>160 or len(data.get('pool',''))>40: raise APIError('Name or pool exceeds its limit.')
                key=(division.id,team.id)
                if db.scalar(select(Entry.id).where(Entry.division_id==division.id,Entry.team_id==team.id)): raise APIError('Team is already entered in this division.')
                data.update(tournament_id=division.tournament_id,color=team.color)
            elif kind=='fixtures':
                stage=scoped(Stage,data.get('stage_id'));a=scoped(Entry,data.get('entry_a_id'));b=scoped(Entry,data.get('entry_b_id'))
                if a.id==b.id or a.division_id!=stage.division_id or b.division_id!=stage.division_id: raise APIError('Two different entries from the stage division are required.')
                planned=datetime.fromisoformat(data.get('planned_at','').replace('Z','+00:00'))
                if planned.tzinfo is None: raise APIError('planned_at requires a timezone offset.')
                tournament=scoped(Tournament,stage.tournament_id)
                local=planned.astimezone(ZoneInfo(tournament.timezone)).date()
                if not tournament.start_date<=local<=tournament.end_date: raise APIError('Fixture date is outside the tournament dates.')
                data['planned_at']=iso(planned);data['planned_duration']=bounded_int(data.get('planned_duration') or 60,60,1,600)
                data.update(tournament_id=stage.tournament_id,division_id=stage.division_id)
                key=(stage.id,a.id,b.id,data['planned_at'])
                if db.scalar(select(Match.id).where(Match.stage_id==stage.id,Match.entry_a_id==a.id,Match.entry_b_id==b.id,Match.planned_at==planned)): raise APIError('Fixture already exists; existing results are never overwritten.')
                if len(data.get('referee_name',''))>120: raise APIError('Referee name exceeds 120 characters.')
            else:
                match=scoped(Match,data.get('match_id'))
                if match.result_status=='finalized' or match.status in {'in_progress','paused','ended'}: raise APIError('Referee metadata can be imported only before play.')
                if not data.get('referee_name') or len(data['referee_name'])>120: raise APIError('Referee name of at most 120 characters is required.')
                key=match.id
            if key in seen: raise APIError('Duplicate row identifier in this import.')
            seen.add(key);validated.append(data)
        except (APIError,ValueError,TypeError) as error:
            errors.append({'row':number,'message':error.message if isinstance(error,APIError) else 'Invalid field or timestamp.'})
    return validated,errors


@bp.get('/imports/template')
@require_org('read')
def import_template():
    kind=request.args.get('type','teams');import_permission(kind)
    output=StringIO();csv.writer(output).writerow(TEMPLATES[kind])
    return send_file(BytesIO(output.getvalue().encode()),mimetype='text/csv',as_attachment=True,download_name='volleyball-'+kind+'-template.csv')


@bp.post('/imports/preview')
@require_org('read')
def import_preview():
    kind=request.form.get('type','teams');import_permission(kind)
    upload=request.files.get('file')
    if not upload or not upload.filename or not upload.filename.lower().endswith('.csv'): raise APIError('Upload a UTF-8 CSV file.')
    raw=upload.read(2*1024*1024+1)
    if len(raw)>2*1024*1024: raise APIError('CSV upload exceeds 2 MB.',413,'file_limit')
    try:
        decoded=raw.decode('utf-8-sig')
        if '\x00' in decoded: raise ValueError('null')
        reader=csv.DictReader(StringIO(decoded),strict=True)
        if reader.fieldnames!=TEMPLATES[kind]: raise APIError('Use the exact columns and ordering from the downloadable template.',details={'expected':TEMPLATES[kind]})
        rows=list(reader)
        if not rows or len(rows)>MAX_IMPORT_ROWS or any(None in row for row in rows): raise APIError('CSV must contain 1–2,000 rows with the expected field count.')
    except (UnicodeDecodeError,csv.Error,ValueError): raise APIError('File must be valid UTF-8 CSV.')
    validated,errors=validate_import(kind,rows)
    job=Job(organization_id=g.organization.id,user_id=g.user.id,kind='import',type=kind,status='review',progress=0,
            data={'rows':rows,'preview':validated[:100],'errors':errors},expires_at=now()+timedelta(hours=24))
    get_db().add(job);get_db().flush();audit('import.previewed','job',job.id,after={'type':kind,'rows':len(rows),'errors':len(errors)});get_db().commit()
    return job_json(job),201


@bp.get('/imports/<job_id>')
@require_org('read')
def get_import(job_id): return job_json(own_job(job_id,'import'))


@bp.post('/imports/<job_id>/commit')
@require_org('read')
def import_commit(job_id):
    db=get_db();job=own_job(job_id,'import')
    job=db.scalar(select(Job).where(Job.id==job.id).with_for_update())
    if job.status!='review': raise APIError('This preview has already been committed.',409,'conflict')
    if job.expires_at<=now(): raise APIError('Preview expired; upload and review again.',410,'expired')
    rows,errors=validate_import(job.type,job.data['rows'])
    if errors:
        job.data={**job.data,'errors':errors};db.commit()
        raise APIError('Correct all row errors before committing. No data was imported.',422,'import_errors',errors)
    for values in rows:
        if job.type=='teams': db.add(Team(organization_id=g.organization.id,name=values['name'],abbreviation=values.get('abbreviation'),color=values['color']))
        elif job.type=='entries': db.add(Entry(organization_id=g.organization.id,division_id=values['division_id'],team_id=values['team_id'],
            tournament_id=values['tournament_id'],name=values['name'],seed=values['seed'],pool=values.get('pool') or None,color=values['color']))
        elif job.type=='fixtures':
            from .competition import schedule_conflicts
            row=Match(organization_id=g.organization.id,tournament_id=values['tournament_id'],division_id=values['division_id'],stage_id=values['stage_id'],
                entry_a_id=values['entry_a_id'],entry_b_id=values['entry_b_id'],planned_at=datetime.fromisoformat(values['planned_at']),
                planned_duration=values['planned_duration'],referee_name=values.get('referee_name') or None)
            db.add(row);db.flush()
            conflicts=schedule_conflicts(row)
            if conflicts: raise APIError('Imported fixture has schedule conflicts. No rows were committed.',422,'schedule_conflicts',conflicts)
        else:
            match=db.scalar(select(Match).where(Match.id==values['match_id']).with_for_update())
            if match.result_status=='finalized' or match.status in {'in_progress','paused','ended'}: raise APIError('Match changed after preview; review again.',409,'conflict')
            match.referee_name=values['referee_name']
    job.status='completed';job.progress=100;job.completed_at=now();job.data={**job.data,'errors':[],'imported_count':len(rows)}
    audit('import.committed','job',job.id,after={'type':job.type,'rows':len(rows),'mode':'all-or-nothing'});db.commit();return job_json(job)


@bp.get('/audit')
@require_org('read')
def audits(): return present_report(build_report('audit',request.args),request.args)


def validate_layout(layout):
    if not isinstance(layout,dict) or len(json.dumps(layout))>30000: raise APIError('Layout must be a JSON object under 30 KB.')
    # Layouts are data consumed by safe DOM code; reject active content and unsafe links.
    def check(value):
        if isinstance(value,str):
            lowered=value.strip().lower()
            if '<' in value or '>' in value or lowered.startswith(('javascript:','data:','vbscript:')): raise APIError('Layout contains unsafe content.')
        elif isinstance(value,dict):
            for key,child in value.items():
                if str(key).lower().startswith('on') or str(key) in {'__proto__','constructor','prototype'}: raise APIError('Unsafe layout key.')
                check(child)
        elif isinstance(value,list):
            if len(value)>100: raise APIError('Layout list is too long.')
            for child in value: check(child)
    check(layout);return layout


@bp.route('/preferences',methods=['GET','POST'])
@require_org('read')
def preferences():
    db=get_db()
    if request.method=='GET':
        query=select(Preference).where(Preference.organization_id==g.organization.id,Preference.user_id==g.user.id)
        if request.args.get('kind'):
            if request.args['kind'] not in {'organizer','report'}: raise APIError('Unknown preference kind.')
            query=query.where(Preference.kind==request.args['kind'])
        rows=list(db.scalars(query))
        return {'items':[serialize(p) for p in rows],'total':len(rows)}
    values=body();kind=values.get('kind','organizer');layout=validate_layout(values.get('layout',{}))
    name=str(values.get('name',layout.get('name','Default'))).strip()
    if kind not in {'organizer','report'} or not name or len(name)>120: raise APIError('Use organizer or report preferences and a name under 120 characters.')
    row=db.scalar(select(Preference).where(Preference.organization_id==g.organization.id,Preference.user_id==g.user.id,Preference.kind==kind,Preference.name==name))
    if not row: row=Preference(organization_id=g.organization.id,user_id=g.user.id,kind=kind,name=name);db.add(row)
    row.layout=layout;db.commit();return serialize(row)


@bp.route('/tournaments/<tournament_id>/display',methods=['GET','PATCH'])
@require_org('read')
def public_display(tournament_id):
    tournament=scoped(Tournament,tournament_id)
    if request.method=='GET': return {'display':tournament.display or {}}
    if g.membership.role not in {'admin','manager'}: raise APIError('Only tournament managers can configure the public display.',403,'forbidden')
    values=body();layout=validate_layout(values.get('display',values.get('layout',{})))
    widgets={'live','standings','upcoming','results','brackets','announcements','weather','sponsors'}
    if 'widgets' in layout and (not isinstance(layout['widgets'],list) or any(not isinstance(w,str) or w not in widgets for w in layout['widgets'])):
        raise APIError('Choose supported display widgets.')
    if 'sizes' in layout and (not isinstance(layout['sizes'],dict) or any(w not in widgets or not isinstance(size,str) or size not in {'normal','wide'} for w,size in layout['sizes'].items())):
        raise APIError('Display widget sizes must be normal or wide.')
    from .admin import image_url
    if 'sponsor_image' in layout:
        layout['sponsor_image']=image_url(layout['sponsor_image'])
    if 'sponsor_url' in layout:
        sponsor_url=image_url(layout['sponsor_url'])
        if sponsor_url and not sponsor_url.lower().startswith('https://'):
            raise APIError('Use a valid https sponsor website URL.')
        layout['sponsor_url']=sponsor_url
    for key,low,high in [('rotation_seconds',5,300),('text_scale',.8,2)]:
        if key in layout:
            try: value=float(layout[key])
            except (ValueError,TypeError):raise APIError(f'{key} must be a number.')
            if not math.isfinite(value) or not low<=value<=high:raise APIError(f'{key} must be between {low} and {high}.')
            layout[key]=value
    if layout.get('background') and (not isinstance(layout['background'],str) or not re.fullmatch('#[0-9a-fA-F]{6}',layout['background'])):raise APIError('Use a #RRGGBB display background.')
    if layout.get('units') and (not isinstance(layout['units'],str) or layout['units'] not in {'celsius','fahrenheit'}):raise APIError('Choose celsius or fahrenheit display weather units.')
    if layout.get('division_id') and scoped(Division,layout['division_id']).tournament_id!=tournament.id:
        raise APIError('Display division must belong to this tournament.')
    if layout.get('court_id'):
        court=scoped(Court,layout['court_id'])
        venues=set((tournament.settings or {}).get('venue_ids',[]))
        venues.update(get_db().scalars(select(TournamentVenue.venue_id).where(TournamentVenue.tournament_id==tournament.id,TournamentVenue.active.is_(True))))
        if court.venue_id not in venues:raise APIError('Display court must belong to a tournament venue.')
    if 'public_reports' in layout and (not isinstance(layout['public_reports'],list) or any(not isinstance(t,str) or t not in PUBLIC_TYPES for t in layout['public_reports'])):
        raise APIError('Choose only supported public report types.')
    previous=tournament.display;tournament.display=layout
    audit('display.changed','tournament',tournament.id,before=previous,after=layout);get_db().commit();return {'display':layout}


@bp.get('/organizations/export')
@require_org('read')
def organization_export():
    if g.membership.role!='admin': raise APIError('Organization administrator access required.',403,'forbidden')
    if permitted_tournaments(): raise APIError('Organization-wide administrator scope is required for business-data export.',403,'forbidden')
    db=get_db();payload={'schema_version':1,'generated_at':iso(now()),'organization':without_secrets(serialize(g.organization)),'purpose':'business-data portability, not infrastructure disaster recovery'}
    for model in (Team,Venue,Court,Tournament,TournamentVenue,Division,Stage,RuleVersion,Entry,Match,CourtAssignment,MatchEvent,ResultRevision,Announcement):
        rows=list(db.scalars(select(model).where(model.organization_id==g.organization.id)))
        values=[]
        for row in rows:
            value=serialize(row)
            # Historical actor IDs remain identifiers; credentials and session details are absent.
            for key in ('session_id','client_action_id','receipt'): value.pop(key,None)
            values.append(without_secrets(value))
        payload[model.__tablename__]=values
    payload['audit_history']=[{key:value for key,value in serialize(a).items() if key not in {'before','after','session_id'}}
        for a in db.scalars(select(Audit).where(Audit.organization_id==g.organization.id))]
    audit('organization.exported','organization',g.organization.id,after={'schema_version':1});db.commit()
    return send_file(BytesIO(json.dumps(payload,ensure_ascii=False,indent=2,default=str).encode()),mimetype='application/json',
                     as_attachment=True,download_name=g.organization.slug+'-business-data.json')
